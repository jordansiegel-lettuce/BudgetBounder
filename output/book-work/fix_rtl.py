from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
f='C:/Users/jordi_vyuby09/Desktop/school/final proj docs/ספר פרויקט גמר - BudgetBounder - מלא.docx'
d=Document(f)
for p in d.paragraphs:
    if p.text.strip() and not re.search('[\u0590-\u05ff]',p.text) and p.alignment!=1:
        p.alignment=0
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if re.search('[\u0590-\u05ff]',r.text):
                        r.font.rtl=True
                        rp=r._r.get_or_add_rPr()
                        for old in list(rp.findall(qn('w:szCs'))):rp.remove(old)
                        sz=OxmlElement('w:szCs');sz.set(qn('w:val'),'18');rp.append(sz)
                        if r.bold:
                            r.font.cs_bold=True
d.save(f)
