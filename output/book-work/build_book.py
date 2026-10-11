from pathlib import Path
import re, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[2]
WORK=Path(__file__).resolve().parent
SRC=Path('C:/Users/jordi_vyuby09/Desktop/school/final proj docs/ספר פרויקט גמר.docx')
OUT=SRC.with_name('ספר פרויקט גמר - BudgetBounder - מלא - עיצוב מקורי.docx')
orig=Document(SRC)
texts=[p.text.strip() for p in orig.paragraphs]
d=Document(SRC)
# Preserve the original cover, logo, forms page, contents pages, styles,
# margins, numbering and first-page order. Replace only the unfinished body.
start=d.paragraphs[113]._p
body=d._element.body
deleting=False
for el in list(body):
    if el is start:deleting=True
    if deleting and el.tag!=qn('w:sectPr'):body.remove(el)
def p(text='',style=None,ltr=False):
    a=d.add_paragraph(style=style or 'Normal')
    # Hebrew follows the same direction and visible order as the source file.
    pp=a._p.get_or_add_pPr(); b=OxmlElement('w:bidi'); b.set(qn('w:val'),'0' if ltr else '1'); pp.append(b)
    if ltr:a.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=a.add_run(text)
    if not ltr:r.font.rtl=True
    return a
def h(text,level=1):
    if level==1:text=re.sub(r'^(\d+)\s+',r'\1. ',text)
    a=p(text,'Normal')
    a.paragraph_format.keep_with_next=True
    a.paragraph_format.space_before=Pt(8 if level==1 else 4)
    for r in a.runs:
        r.bold=True;r.font.cs_bold=True;r.underline=True;r.font.size=Pt(14)
    return a
def paras(*items):
    for t in items:p(t)
def bullets(items):
    for t in items:p('• '+t,'List Paragraph')
def table(headers,rows,widths=None):
    if d.paragraphs:d.paragraphs[-1].paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    pr=t._tbl.tblPr
    bd=OxmlElement('w:bidiVisual'); pr.append(bd)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        for i,v in enumerate(row):t.add_row() if False else None
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            tcpr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');tcpr.append(shade)
            for a in c.paragraphs:
                a.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                b=OxmlElement('w:bidi');a._p.get_or_add_pPr().append(b)
                a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1.05
                if len(t.rows)<=10:a.paragraph_format.keep_with_next=ri<len(t.rows)-1
                for r in a.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ri==0
                    rp=r._r.get_or_add_rPr();fonts=rp.get_or_add_rFonts();fonts.set(qn('w:cs'),'David')
    p()
    return t
def old(a,b):
    for i in range(a,b+1):
        if texts[i]:p(texts[i])
def diagram(name,kind):
    im=Image.new('RGB',(1600,1000),'white');dr=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28)
    small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',23)
    def box(x,y,w,hh,text):
        dr.rounded_rectangle((x,y,x+w,y+hh),radius=12,fill='#eef2f6',outline='#273f55',width=3)
        lines=text.split('\n')
        for i,line in enumerate(lines):dr.text((x+w/2,y+hh/2+(i-(len(lines)-1)/2)*35),line,font=font,fill='#182a3a',anchor='mm')
    def arrow(a,b,label=''):
        import math
        dr.line([a,b],fill='#34495e',width=3);ang=math.atan2(b[1]-a[1],b[0]-a[0])
        pts=[b,(b[0]-18*math.cos(ang-.5),b[1]-18*math.sin(ang-.5)),(b[0]-18*math.cos(ang+.5),b[1]-18*math.sin(ang+.5))];dr.polygon(pts,fill='#34495e')
        if label:dr.text(((a[0]+b[0])/2,(a[1]+b[1])/2-22),label,font=small,fill='#182a3a',anchor='mm')
    if kind=='arch':
        box(80,100,520,140,'Mobile app\nReact Native + Expo');box(1000,100,520,140,'Admin website\nReact + Vite')
        box(540,420,520,150,'ASP.NET Core API\nAuth + business logic')
        box(120,780,500,140,'SQL Server\nEF Core persistence');box(1000,780,480,140,'Groq API\nAI coach + mission drafts')
        arrow((400,240),(680,420),'JSON / HTTPS');arrow((1220,240),(920,420),'JSON / HTTPS');arrow((650,570),(420,780),'EF Core');arrow((1000,570),(1230,780),'Server API key')
    elif kind=='use':
        box(30,430,240,100,'User');box(1330,430,240,100,'Admin')
        left=['Register / sign in','Record transactions','Budget / savings goals','Missions / rewards','Play Tower / submit result','Ask AI coach']
        right=['Users / account status','Review AI mission drafts','Create system missions','Analytics / CSV export','Game results / rewards','Audit / monitoring']
        for i,txt in enumerate(left):box(340,40+i*155,430,95,txt);arrow((270,480),(340,87+i*155))
        for i,txt in enumerate(right):box(830,40+i*155,430,95,txt);arrow((1330,480),(1260,87+i*155))
    elif kind=='seq':
        xs=[160,540,960,1400]
        for x,txt in zip(xs,['User','Mobile','API','SQL Server']):
            box(x-130,30,260,85,txt);dr.line([(x,115),(x,940)],fill='#9ba8b4',width=2)
        steps=[(0,1,'Finish Tower',200),(1,2,'POST game result + JWT + result ID',310),(2,3,'Find existing user / result ID',420),(3,2,'Existing result or no match',520),(2,2,'Validate limits; calculate XP',620),(2,3,'Save result + XP + level history',730),(3,2,'Commit',820),(2,1,'Stored result / awarded XP',910)]
        for a,b,txt,y in steps:
            if a==b:box(xs[a]-210,y-35,420,70,txt)
            else:arrow((xs[a],y),(xs[b],y),txt)
    else:
        box(30,60,300,110,'User / Mobile');box(1230,60,340,110,'Administrator')
        box(40,390,430,120,'1. Finance operations');box(585,390,430,120,'2. Missions / progression');box(1110,390,430,120,'3. Admin reporting')
        box(400,800,800,125,'Data stores\nUsers | Transactions | Goals | Missions | GameSessions')
        box(570,50,460,120,'4. AI request via server\nGroq provider')
        arrow((180,170),(220,390),'Entries / goals');arrow((330,120),(570,120),'Question');arrow((850,170),(800,390),'Mission drafts');arrow((1390,170),(1340,390),'Filters / approvals')
        arrow((270,510),(530,800),'Read / write');arrow((800,510),(800,800),'Progress / XP');arrow((1290,510),(1090,800),'Read / audit')
    path=WORK/(name+'.png');im.save(path)
    a=d.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.add_run().add_picture(str(path),width=Cm(16.5))
    a.paragraph_format.keep_with_next=True

# Completed body in the same chapter order as the original contents pages.
h('1 שם הפרויקט')
p('BudgetBounder היא מערכת לניהול תקציב אישי ולעידוד הרגלים כלכליים באמצעות חוויית שימוש משחקית. הפרויקט כולל אפליקציית מובייל למשתמש, אתר ניהול ושרת API משותף.')
h('2 רקע')
h('2.1 תיאור ורקע כללי',2)
old(117,117)
p('המשתמש מזין הכנסות והוצאות, מגדיר תקציב ויעדי חיסכון ועוקב אחר התקדמותו. ביצוע פעולות ומשימות מזכה בנקודות ניסיון, והמשחק משתלב בתהליך ההתמדה. מנהל המערכת מפקח על משתמשים, משימות, המלצות ותוצאות משחק מתוך אתר ייעודי.')
h('2.2 מטרות המערכת',2);old(119,124)
h('3 המצב בשוק והבעיות')
h('3.1 סקירת מצב קיים בשוק',2)
p('כלים לניהול תקציב מרכזים נתונים כספיים, מציגים חלוקה לקטגוריות ומאפשרים לעקוב אחר יעדים. עצם הצגת הנתונים אינה מבטיחה שהמשתמש יתמיד ברישום ויהפוך את המידע לפעולות יומיומיות. BudgetBounder מתמקדת בחיבור בין המעקב לבין משימות והתקדמות משחקית.')
p('Mint, שהוזכרה באפיון המקורי, משמשת כאן דוגמה היסטורית. Intuit הודיעה על שילובה ב־Credit Karma ועל העברת משתמשיה; לכן אין להציג אותה כאפליקציה עצמאית פעילה להשוואת מוצרים בשנת ההגשה. המקור להודעה מופיע בפרק 12.5.')
p('ההשוואה הרלוונטית היא בין איסוף אוטומטי של נתונים, שהיה חלק מחוויית Mint, לבין רישום ידני ב־BudgetBounder; ובין התמקדות בדוחות כספיים לבין שילוב דוחות, משימות, XP ומשחק. בפרויקט לא נמצא חיבור ישיר לחשבונות בנק.')
h('3.2 בעיות במצב הקיים',2)
bullets(['קושי להתמיד בהזנת נתונים לאורך זמן כאשר המשתמש אינו מקבל משוב ברור על הפעולות שביצע.','קושי לתרגם תרשים הוצאות לצעד מעשי, כגון שינוי הרגל או קביעת יעד חיסכון.','עומס מידע עבור משתמשים ללא ניסיון בניהול תקציב.','פער בין הצבת יעד לבין מעקב שבועי וחודשי אחר ההתקדמות.'])
p('אלו בעיות תכנון שהפרויקט מבקש לתת להן מענה. לא נערך במסגרת הספר מחקר משתמשים שמוכיח שיעור נטישה או יתרון כמותי על כל המוצרים בשוק.')
h('4 מה הפרויקט אמור לחדש או לשפר');old(154,154)
p('השיפור המוצע הוא מעגל פעולה קצר: המשתמש מתעד פעולה כספית, רואה את השפעתה על התקציב ועל המשימות ומקבל משוב על התקדמותו. משימות AI עוברות אישור מנהל לפני פרסום, ותוצאות המשחק מאומתות בשרת לפני הענקת XP. כך הרכיב המשחקי נשען על נתונים ותהליכים מוגדרים.')
h('5 דרישות מערכת ופונקציונאליות')
h('5.1 דרישות מערכת',2)
table(['רכיב','דרישה'],[
('מובייל','מכשיר Android או iOS התואם לגרסת Expo שבה נבנה הפרויקט; הגרסאות המדויקות נקבעות לפי תוצר הבנייה. נדרש חיבור לרשת לפעולות מול השרת.'),
('מחשב פיתוח','Windows או סביבה מתאימה אחרת, .NET SDK 10, Node.js התואם לתלויות, Git ועורך קוד.'),
('שרת','ASP.NET Core על .NET 10, גישה ל־SQL Server ותצורת JWT. לשירותי AI נדרש מפתח Groq בצד השרת.'),
('מנהל','דפדפן עדכני התומך ביישום React וחשבון בעל תפקיד Admin.'),
('הרשאות מכשיר','תמונות או מצלמה לצירוף קבלה, מיקום לציון בית עסק והתראות לתזכורות; יש לטפל גם בסירוב הרשאה.'),
('אבטחה ותפעול','אימות JWT, בדיקת בעלות על נתונים והרשאות מנהל. בפריסה נדרשים HTTPS, ניהול סודות וגיבויים.')],[3,14])
p('המשחק הנוכחי הוא Tower המובנה באפליקציה. דרישות Unity WebGL, שרת קובצי Unity ו־WebView מהאפיון המוקדם אינן נדרשות למימוש זה. אין בקוד שירות Python נפרד ל־AI; קריאות הספק מתבצעות מתוך שרת ה־API.')
h('5.2 דרישות פונקציונאליות',2)
for title,body in [
('הרשמה והתחברות','יצירת חשבון באמצעות שם, דואר אלקטרוני וסיסמה, התחברות, זיהוי המשתמש ושליפת פרופיל. משתמש רגיל אינו מקבל תפקיד מנהל בעת ההרשמה.'),
('ניהול הכנסות והוצאות','הזנת תנועה ידנית, בחירת סוג, סכום, תאריך וקטגוריה; סיווג אוטומטי מבוסס מילות מפתח כאשר נבחר Auto; צפייה בהיסטוריה ובמגמות. ניתן לצרף פרטי קבלה ומיקום.'),
('תקציב וחיסכון','שמירת תקציב חודשי, הצגת הוצאות ויתרה, יצירת יעד חיסכון ורישום הפקדות. הצגת התקדמות שבועית וחודשית מתבססת על הפקדות מתוארכות.'),
('משימות והתקדמות','הצגת משימות זמינות, בדיקת התקדמות, הענקת XP לאחר עמידה בתנאים, עדכון רמות וחשיפת הישגים. משימות AI נוצרות כטיוטות ומחייבות אישור מנהל.'),
('מאמן פיננסי','שליחת שאלה לשרת וקבלת תשובה המבוססת על תנועות אחרונות ויעדי חיסכון. תכנית החיסכון היא המלצה טקסטואלית; היא אינה מעבירה כסף ואינה משנה נתונים ללא פעולה מתאימה.'),
('משחק Tower','הפעלת המשחק מהמובייל, פתיחת קומות בהתאם להתקדמות, שליחת תוצאה ומזהה ייחודי לשרת, בדיקת תקינות וחישוב XP. מצב הדגמה לפיתוח אינו מעניק XP לחשבון.'),
('ניהול משתמשים','רשימת משתמשים, חיפוש וסינון, צפייה בנתוני פעילות והקפאת חשבון. חשבון מוקפא נדחה גם כאשר מוצג אסימון שהונפק לפני ההקפאה.'),
('ניהול תוכן ומעקב','יצירת משימות, עריכה ואישור טיוטות, ניהול הגדרות תגמולים, בדיקת תוצאות משחק, יומן ביקורת ודוחות מסוננים עם ייצוא CSV.')]:
    h(title,3);p(body)
h('6 בעיות צפויות במהלך הפיתוח ופתרונות')
h('6.1 תיאור הבעיות',2)
paras('המערכת משלבת מובייל, אתר ניהול, בסיס נתונים וספק AI. תקלה באחד הרכיבים עלולה להשאיר את המשתמש בלי תשובה, ליצור חוסר התאמה בתצוגה או לגרום לשליחה חוזרת של פעולה.',
'סכומים, גבולות חודש ותאריכי פעילות מחייבים הגדרה עקבית. השלמת משימה או העלאת תוצאת משחק במקביל מכמה בקשות עלולות לגרום להענקת XP כפולה. בנוסף, תוכן שנוצר באמצעות AI עלול להיות לא תקין או לא מתאים למשימה מדידה.',
'בהרצה על מכשיר פיזי יש להתמודד עם כתובת שרת נגישה, תאימות Expo, הרשאות והתנהגות שונה בין Android ל־iOS. גדילת טבלאות התנועות והמשחקים עלולה להאט דוחות ושליפות.')
h('6.2 פתרונות אפשריים',2)
table(['בעיה','פתרון ומצבו'],[
('שליחה חוזרת','מזהה ClientResultId ואינדקס ייחודי למשתמש ולתוצאה; השרת מחזיר תוצאה קיימת במקום להעניק שוב XP.'),
('עדכון משימה במקביל','בדיקות מצב והתקדמות ושדות ConcurrencyCheck למניעת קבלת שינוי המבוסס על מצב ישן.'),
('AI לא זמין או לא תקין','בדיקת תצורה ותשובת ספק; יצירת משימות כטיוטות ואישור מנהל. כשל AI אינו אמור למחוק תנועות או יעדים.'),
('חישובי זמן','חלון חודש מתחיל בתחילת החודש ב־UTC ומסתיים לפני תחילת החודש הבא; בדיקות גבול ייעודיות.'),
('נתונים רבים','אינדקסים לפי משתמש ותאריך, סינון ודפדוף בממשק ניהול. בדיקות עומס והגבלת נפח קבלות הן עבודת תפעול להמשך.'),
('רשת והרשאות','חיווי טעינה ושגיאה, בדיקת כתובת API וטיפול בסירוב הרשאה. עבור תוצאות משחק קיים מנגנון ניסיון חוזר; אין להניח שכל פעולה עובדת ללא רשת.')],[4,13])
h('7 פתרון טכנולוגי נבחר')
h('7.1 טופולוגית הפתרון',2)
p('הפתרון בנוי במודל לקוח ושרת. המובייל ואתר הניהול שולחים בקשות JSON לשרת ASP.NET Core. השרת מאמת את המשתמש, מפעיל לוגיקה עסקית וניגש ל־SQL Server באמצעות Entity Framework Core. רק השרת פונה לספק Groq עם מפתח השירות.')
diagram('architecture','arch');p('איור 1 מבנה הרכיבים וקשרי התקשורת. סביבת הייצור צריכה להגן על התעבורה באמצעות HTTPS.')
h('7.2 טכנולוגיות בשימוש',2)
table(['תחום','טכנולוגיה','תפקיד'],[('אפליקציה','React Native ו־Expo','מסכי מובייל וגישה לשירותי המכשיר'),('ניהול','React ו־Vite','ממשק מנהל בדפדפן'),('שרת','ASP.NET Core 10','API, אימות וביצוע פעולות'),('נתונים','SQL Server ו־EF Core','מיפוי ישויות, שאילתות ומיגרציות'),('AI','Groq Chat Completions','יצירת תשובות וטיוטות משימות'),('בדיקות','xUnit, Jest, Vitest','בדיקות שרת, מובייל ואתר')],[3,5,9])
h('7.3 שפות הפיתוח',2)
paras('TypeScript משמשת במסכי המובייל, באתר הניהול ובפונקציות עזר. רכיבי TSX מחברים בין מבנה המסך לבין לוגיקת התצוגה. CSS מגדיר את עיצוב האתר.',
'C# משמשת לבקרים, למודלים, לשירותים ולבדיקות השרת. SQL משמשת בבסיס הנתונים ובסקריפטים של המיגרציות. JSON הוא פורמט החלפת הנתונים והגדרות מסוימות; אינו שפת תכנות של רכיב עצמאי.')
h('7.4 תיאור הארכיטקטורה הנבחרת',2)
paras('שכבת התצוגה מנהלת ניווט, טפסים ומצבי טעינה. שכבת ה־API מכילה Controllers שמקבלים בקשות ומחזירים קודי HTTP. שירותים משותפים מרכזים חישובי תקציב, סיווג תנועות, XP ותוצאות משחק. BudgetBounderDbContext ממפה את הישויות ושומר אותן במסד הנתונים.',
'השרת הוא מקור הסמכות להרשאות ולהענקת XP. לקוח יכול להציג התקדמות ולשלוח בקשה, אך אינו רשאי לקבוע בעצמו את זהות בעל הנתונים או את הפרס. ההפרדה מאפשרת להשתמש באותם כללים מהאפליקציה ומאתר הניהול ולבדוק לוגיקה ללא תלות במסך.')
h('7.5 חלוקה לתכניות ומודולים',2)
for name,desc in [('משתמשים והרשאות','UsersController ו־CurrentUserService מטפלים בזיהוי משתמש, פרופיל ותפקיד. Program מגדיר אימות JWT ומדיניות AdminOnly.'),('תנועות וסיכומים','TransactionsController, TransactionCategorizer ו־FinancialSummaryService אחראים לתנועות, קטגוריות וסיכומים חודשיים.'),('תקציב וחיסכון','BudgetsController ו־SavingGoalController מנהלים תקציב, יעדים ורישום הפקדות מתוארכות.'),('משימות והתקדמות','MissionsController, StaticMissionService ו־ProgressionService מנהלים משימות, XP, רמות ורצף פעילות.'),('AI','AiController מטפל במאמן; מנגנון יצירת המשימות משתמש בספק AI. מסכי המנהל מאפשרים בקרת תוכן ואישור טיוטות.'),('משחק ותגמולים','לוגיקת Tower במובייל ו־GameSessionsController בשרת מנהלים תוצאות; GameResultValidator בודק ערכים; AchievementsController מטפל בהישגים.'),('ניהול ודוחות','בקרי Admin ומסכי האתר מספקים משתמשים, משימות, אנליטיקה, תגמולים, יומן ביקורת וניטור.')]:h(name,3);p(desc)
h('7.6 סביבת השרת',2)
paras('השרת מכוון ל־net10.0 ומשתמש ב־SQL Server דרך ספק EF Core. חיבור המסד, פרטי JWT ומפתח Groq נקראים מתצורה. יש לשמור ערכים סודיים מחוץ לקוד ולספר הפרויקט.',
'בסביבת פיתוח קיימים Swagger וכתובת API מקומית. לקוח מובייל על מכשיר פיזי נדרש לפנות לכתובת מחשב נגישה ברשת ולא ל־localhost של המכשיר. אתר הניהול משתמש ב־VITE_API_URL והמובייל ב־EXPO_PUBLIC_API_URL.',
'לפני פריסה יש להחיל מיגרציות על המסד המיועד, להגדיר כתובות CORS מתאימות, HTTPS וגיבוי. מסמכי ההרצה במאגר מתארים עדכון מקומי ואינם מהווים אישור שהגרסה נפרסה בשרת מרוחק.')
h('7.7 ממשק המשתמש והלקוח',2)
table(['מסך','תוכן ופעולות'],[('כניסה והרשמה','זיהוי משתמש והצגת שגיאה במקרה של פרטים לא תקינים'),('בית ופרופיל','תקציב, XP, רמה, רצף פעילות והישגים'),('תנועות','רשימה, מגמות והוספת הכנסה או הוצאה'),('יעדי חיסכון','יצירת יעד, סכום יעד, תאריך והוספת הפקדה'),('משימות','משימות זמינות, מצב התקדמות ובקשת יצירת טיוטות'),('מאמן AI','שאלה חופשית ובקשת תכנית חיסכון'),('Tower','בחירת קומה, משחק ותוצאה'),('אתר מנהל','ניווט בין משתמשים, משימות, אנליטיקה, משחקים ויומן ביקורת')],[4,13])
p('העיצוב משלב רקעים כהים, פעולות מודגשות וגוונים כתומים וזהובים. נתונים מספריים מוצגים לצד פעולות קצרות וברורות. מסכי הניהול כוללים סינון ודפדוף, ומסכי הקלט צריכים להציג הודעה כאשר השמירה נכשלת.')
shot=ROOT/'BudgetBounder/BudgetBounder.Web/BudgetBounder-web/output/playwright/admin-missions-desktop.png'
if shot.exists():
    d.add_picture(str(shot),width=Cm(16.5));p('איור 2 מסך ניהול המשימות מתוך תיעוד הבדיקות במאגר. התמונה נוצרה עם נתוני בדיקה ולא עם נתוני משתמשים בסביבת ייצור.')
h('7.8 ממשקים למערכות אחרות ו־API',2)
table(['ממשק','קלט ופלט','ייעוד'],[('POST /api/users/register','פרטי הרשמה; פרופיל חדש','פתיחת חשבון'),('POST /api/users/login','דואר וסיסמה; JWT ופרופיל','התחברות'),('GET /api/users/me','JWT; פרופיל נוכחי','רענון זהות ונתוני משתמש'),('POST /api/ai/chat','שאלה; תשובת מאמן','השרת מצרף הקשר כספי ופונה ל־Groq'),('Groq API','messages ומודל; תוכן תשובה','שירות חיצוני הנקרא מהשרת בלבד'),('Expo APIs','הרשאת מכשיר; תמונה, מיקום או התראה','שירותי מכשיר לפי פעולה')],[5,6,6])
p('בקשות אישיות כוללות Authorization: Bearer ואסימון JWT. השרת מבחין בין בקשה שגויה, משתמש לא מזוהה, היעדר הרשאה ומשאב שלא נמצא. חוזי שאר הבקרים זמינים בקוד וב־Swagger בסביבת פיתוח.')
h('7.9 שימוש בחבילות תוכנה',2)
table(['חבילה','שימוש'],[('expo-router','ניווט במובייל לפי קבצים'),('expo-secure-store','אחסון מאובטח של נתוני התחברות במכשירים נתמכים'),('expo-image-picker / expo-location','בחירת תמונה וגישה למיקום'),('expo-notifications','תזכורות מקומיות והגדרות התראות'),('axios','שליחת בקשות HTTP בצד הלקוח'),('react-router-dom','ניווט באתר המנהל'),('Microsoft.EntityFrameworkCore.SqlServer','גישה למסד הנתונים'),('JwtBearer / BCrypt.Net-Next','אימות אסימונים וגיבוב סיסמאות'),('Swashbuckle.AspNetCore','תיעוד API בפיתוח')],[7,10])
h('8 שימוש במבני נתונים וארגון קבצים')
h('8.1 מבני הנתונים',2)
p('מודל הנתונים בפועל משתמש ב־Transactions במקום Expenses, ב־Missions במקום Tasks וב־GameSessions במקום GameResults. הישגים ממומשים באמצעות RewardDefinitions ו־UserRewards. מילון הנתונים להלן מחליף את שמות הטבלאות מהתכנון המוקדם ומתבסס על מחלקות Models הנוכחיות.')
p('טיפוסים מוצגים כפי שהם בקוד C#. הסימן ? מציין ערך שעשוי להיות חסר. קשר לוגי לפי UserId אינו מעיד בהכרח על אילוץ מפתח זר פיזי בכל טבלה; המיגרציות הן המקור לסכמת SQL המדויקת. סכומים כספיים ממומשים כעת ב־double; מעבר ל־decimal הוא שיפור אפשרי לדיוק חישובים.')
descs={'User':'זהות, הרשאות והתקדמות משתמש','Transaction':'הכנסה או הוצאה ופרטי קבלה','MonthlyBudget':'תקציב לחודש מסוים','SavingGoal':'יעד חיסכון ומצב התקדמות','SavingContribution':'הפקדה מתוארכת ליעד','Mission':'משימה מדידה ותהליך אישור','GameSession':'תוצאת משחק ומצב אימות','RewardDefinition':'הגדרה של תגמול או הישג','UserReward':'תגמול שנפתח למשתמש','LevelHistory':'צילום מצב XP ורמה לאורך זמן','ReminderPreferences':'העדפות תזכורות','AiRecommendation':'המלצה ותהליך סקירת מנהל','AdminAuditLog':'תיעוד פעולת מנהל'}
fielddesc={'Id':'מזהה רשומה','UserId':'מזהה משתמש','FullName':'שם מלא','Email':'כתובת דואר','PasswordHash':'גיבוב סיסמה','Level':'רמה','XP':'נקודות ניסיון','Role':'תפקיד','IsActive':'מצב חשבון או הגדרה','CreatedAt':'מועד יצירה','LastActiveAt':'מועד פעילות שנרשם','CurrentStreak':'רצף נוכחי','LongestStreak':'רצף מרבי','LastActivityDate':'יום פעילות אחרון','Title':'כותרת','Amount':'סכום','Type':'סוג תנועה','Category':'קטגוריה','Date':'תאריך תנועה','ReceiptImageDataUrl':'תמונה כ־Data URL','Latitude':'קו רוחב','Longitude':'קו אורך','MerchantAddress':'כתובת בית עסק','Month':'חודש התקציב','UpdatedAt':'מועד עדכון','TargetAmount':'סכום יעד','CurrentAmount':'סכום נצבר','Deadline':'מועד יעד','IsCompleted':'הושלם','SavingGoalId':'מזהה יעד','Description':'תיאור','XPReward':'פרס XP','Difficulty':'דרגת קושי','CompletedAt':'מועד השלמה','ExpiresAt':'תפוגה','MissionType':'סוג משימה','TargetValue':'ערך נדרש','CurrentProgress':'התקדמות נוכחית','IsAiGenerated':'נוצר באמצעות AI','ReviewStatus':'מצב אישור','ReviewedByAdminId':'מנהל סוקר','ReviewedAt':'מועד סקירה','ClientResultId':'מזהה תוצאה מהלקוח','Score':'ניקוד','Coins':'מטבעות','SavingsStars':'כוכבי חיסכון','DurationSeconds':'משך בשניות','AwardedXp':'XP שהוענק','ValidationState':'תוצאת אימות','ValidationReason':'סיבת סימון','SubmittedAt':'מועד שליחה','Code':'קוד תגמול','Name':'שם','CosmeticType':'סוג חזותי','RewardDefinitionId':'מזהה הגדרת תגמול','UnlockedAt':'מועד פתיחה','RecordedAt':'מועד תיעוד','Content':'תוכן המלצה','Status':'מצב','AdminUserId':'מזהה מנהל','Action':'סוג פעולה','TargetType':'סוג יעד','TargetId':'מזהה יעד פעולה','Details':'פרטי פעולה','DailyLogEnabled':'תזכורת יומית פעילה','DailyLogHour':'שעת תזכורת','BudgetAlertsEnabled':'התראות תקציב','BudgetThresholdPercent':'סף באחוזים','GoalRemindersEnabled':'תזכורת יעד','GoalReminderWeekday':'יום תזכורת יעד','GoalReminderHour':'שעת תזכורת יעד','MissionAlertsEnabled':'התראות משימות'}
for name,desc in descs.items():
    h(name,3);p(desc)
    code=(ROOT/f'BudgetBounder/BudgetBounder.API/BudgetBounder.Api/Models/{name}.cs').read_text(encoding='utf-8-sig')
    fields=re.findall(r'public\s+(int\??|double\??|string\??|bool|DateTime\??|DateOnly\??|Guid|TransactionType)\s+(\w+)\s*\{\s*get;',code)
    table(['שדה','טיפוס C#','משמעות'],[(f,t,fielddesc.get(f,f)) for t,f in fields],[6,4,7])
p('קשרים מרכזיים: משתמש מחזיק תנועות, תקציבים, יעדים, משימות ותוצאות משחק. הפקדה מפנה ליעד חיסכון; UserReward מחבר משתמש להגדרת תגמול. AdminAuditLog מתעד את המנהל שביצע את הפעולה. הרשומות ההיסטוריות מאפשרות להסביר כיצד השתנתה ההתקדמות.')
p('אינדקסים ייחודיים מונעים יותר מתקציב אחד למשתמש ולחודש, כפילות בתוצאת משחק לפי משתמש ומזהה לקוח, וכפילות בפתיחת אותו תגמול. אינדקסים לפי משתמש ותאריך תומכים בשליפת תנועות, הפקדות והיסטוריית רמות.')
h('8.2 שיטת האחסון',2)
paras('SQL Server הוא האחסון המתמשך. EF Core ממפה אובייקטים לטבלאות ומנהל שינויים בסכמה באמצעות Migrations. רשימות בזיכרון משמשות לעיבוד תנועות, ומילון לפי קטגוריה משמש להחזרת סיכומי הוצאות.',
'בתיקיית השרת נמצאים Controllers, Models, Dtos, Services, Authorization, Data ו־Migrations. במובייל נמצאות תיקיות app למסכים ו־src ללוגיקה משותפת; באתר נמצאות pages/admin, components, services ו־context. הפרדה זו מסייעת לאתר את האחריות לכל פעולה.',
'תמונת קבלה נשמרת בשדה ReceiptImageDataUrl של התנועה במימוש הקיים. אין להסיק מכך שקיים אחסון קבצים חיצוני או זיהוי OCR. היסטוריית הפקדות ורמות מתועדת רק מהמועד שבו הוחל המנגנון; אין לשחזר מועדים ישנים בהשערה.')
h('8.3 מנגנוני התאוששות',2)
paras('במימוש קיים טיפול בשליחה חוזרת של תוצאות משחק: מזהה התוצאה נשמר, והשרת מזהה כפילות. שינויים תלויים כגון שמירת תוצאה ועדכון XP מתבצעים יחד במסגרת פעולת השמירה המתאימה. בדיקות התנגשות במשימות מונעות עדכון על סמך מצב ישן.',
'תכנית התאוששות לפריסה: לבצע גיבוי מסד לפני מיגרציה, לשמור עותק מחוץ לשרת, לבדוק שחזור בסביבה נפרדת ולתעד את גרסת הקוד ואת גרסת הסכמה. במקרה של כשל יש לעצור כתיבות אם נדרש, לשחזר עותק מאומת ולהריץ בדיקות כניסה, תנועות ויעדים.',
'בקוד שנבדק לא תועדו גיבוי אוטומטי, אשכול שרתים או מעבר אוטומטי לשרת חלופי. אלה דרישות תפעול להמשך, ולא יכולות שנבדקו. יעדי זמן התאוששות ואובדן מידע מרבי ייקבעו בהתאם לתשתית האירוח.')
h('9 תרשימי מערכת מרכזיים')
p('התרשימים מציגים את תהליכי BudgetBounder. התוויות בתוך התרשימים באנגלית כדי להתאים לשמות הרכיבים בקוד; ההסבר מופיע בעברית.')
h('9.1 תרשים מקרי שימוש',2);diagram('usecases','use')
p('איור 3 מקרי השימוש העיקריים של משתמש ומנהל. פעולות אישיות דורשות התחברות; פעולות ניהול דורשות Admin. ספק AI הוא שירות חיצוני המופעל על ידי השרת, ואינו משתמש בעל גישה למסד.')
h('9.2 תרשים רצף',2);diagram('sequence','seq')
p('איור 4 תהליך שליחת תוצאת Tower. אם נמצאה תוצאה קודמת עם אותו מזהה, מוחזרת התוצאה השמורה בלי פרס נוסף. אם הערכים אינם בטווח התקין, התוצאה מסומנת Flagged ומתקבל אפס XP. המסלול המוצג לשמירה ולצבירת XP הוא המסלול התקין.')
h('9.3 תרשים זרימת נתונים',2);diagram('dataflow','dfd')
p('איור 5 תהליכי כספים, משימות, AI וניהול מול מאגרי הנתונים. הנתונים הנשלחים לספק AI נבחרים בשרת; טיוטת משימה דורשת סקירת מנהל. דוחות נגזרים מהנתונים שנשמרו ואינם משנים תנועות כספיות.')
h('10 תיאור המרכיב האלגוריתמי והחישובי')
h('10.1 הבעיה שבא לפתור',2);old(438,442)
p('המרכיב החישובי ממיר תנועות ופעולות למדדים ברורים: הוצאות חודשיות, יתרת תקציב, התקדמות חיסכון, XP ורמה. כללי החישוב הדטרמיניסטיים מבוצעים בשרת; AI משמש לניסוח המלצות וטיוטות ולא לחישוב יתרה מוסמכת.')
h('10.2 איסוף מידע וניתוחים סטטיסטיים',2)
paras('המידע נאסף מהזנת תנועות, הפקדות, השלמת משימות ותוצאות משחק. דוחות המנהל מציגים פעילות, השלמות, הענקות XP ותצפיות רמה. LastActiveAt ונתוני פעילות אינם מדידה מלאה של זמן שהייה בכל מסך.',
'בחישוב חודש מסננים לפי UserId ולפי טווח [תחילת החודש, תחילת החודש הבא). מסכמים Income בנפרד מ־Expense ומקבצים הוצאות לפי Category. יתרת התקציב היא התקציב שהוגדר פחות ההוצאות; היא אינה יתרת חשבון בנק.',
'דוגמה חישובית: עבור תקציב 3,000 ש״ח והוצאות 1,200 ש״ח, יתרת התקציב היא 1,800 ש״ח ושיעור הניצול הוא 40%. עבור יעד 5,000 ש״ח וחיסכון נצבר 1,250 ש״ח, ההתקדמות היא 25%. אלה ערכי הדגמה, לא נתוני משתמש אמיתי.',
'סיווג אוטומטי מתבצע לפי כללי מילות מפתח בעברית ובאנגלית. התאמה ראשונה קובעת קטגוריה, וכאשר לא נמצאה התאמה מוחזר Other. מדובר במסווג מבוסס כללים; למשתמש נשמרת האפשרות לבחור קטגוריה ידנית.')
table(['רמה','סף XP'],[(str(i+1),str(x)) for i,x in enumerate([0,100,250,500,1000,2000,3000,4000,4500,5000])],[8,9])
p('חישוב פרס משחק: מחלקים ניקוד ב־20 בחלוקה של מספרים שלמים, מוסיפים מטבעות ועשר נקודות לכל כוכב חיסכון, ומגבילים את התוצאה לטווח 1 עד 500. תוצאה שסומנה כלא תקינה מקבלת אפס. לדוגמה: ניקוד 1,000, עשרה מטבעות ושני כוכבים מזכים ב־80 XP.')
p('רצף פעילות אינו גדל מפעולה שנייה באותו יום. פעילות ביום העוקב מגדילה אותו באחד; לאחר הפסקה הוא מתחיל מחדש באחד. המערכת שומרת גם את הרצף הארוך ביותר. חישוב הרמה משתמש בספי XP קבועים, ולא ברמה שמסר הלקוח.')
p('בקשת המאמן כוללת עד 20 תנועות אחרונות ונתוני יעדים. יצירת משימות משתמשת בתנועות אחרונות ובהוראות ליצירת משימה מדידה. אין במימוש שנבדק מנגנון קהילתי לסינון משתמשים לפי מרחק או מעקב כללי אחר משך שהייה בכל מסך.')
h('11 אבטחת מידע')
paras('סיסמאות נשמרות כגיבוב BCrypt ונבדקות באמצעות Verify; הן אינן נשמרות כסיסמה גלויה ואין מדובר בהצפנה שנועדה לפענוח. ההתחברות מחזירה JWT חתום, ובאימות נבדקים מנפיק, קהל יעד, חתימה ותוקף.',
'זהות המשתמש נלקחת מהאסימון בצד השרת. לפני פעולות על משאב אישי יש לבדוק בעלות; פעולות ניהול משתמשות במדיניות AdminOnly. בעת אימות אסימון נבדק שהמשתמש עדיין קיים ושהחשבון פעיל. כך הקפאת חשבון חלה גם על בקשות מאסימון קיים.',
'הגנה על הלוגיקה העסקית כוללת אימות סכומים וטווחים, בדיקת התקדמות משימה, בקרת תוצאות משחק ואינדקסים למניעת כפילות. יומן מנהל מתעד פעולות לצורך ביקורת. בדיקת טווחי ניקוד אינה מוכיחה שכל תוצאת משחק אמיתית, ולכן אין להציג אותה כמנגנון מלא למניעת רמאות.',
'בפריסה יש להשתמש ב־HTTPS, לשמור מפתחות ומחרוזות חיבור במנגנון סודות, להגביל CORS ולצמצם את הרשאות חשבון המסד. אין להכניס מפתח API או פרטי התחברות למסמך, למאגר ציבורי או לתוצרי לקוח.',
'תיאורי עסקאות, סכומים ויעדים עשויים להיכלל בהקשר הנשלח ל־Groq. יש ליידע את המשתמש ולהעביר רק מידע הדרוש לפעולה. בקוד המאמן קיימת הוראה להתייחס לתיאורי תנועות כנתונים ולא כהוראות; אין בכך הבטחה מוחלטת נגד הזרקת הנחיות.',
'הצפנת מסד במנוחה, הגבלת קצב כניסה, אימות רב־שלבי, מדיניות מחיקת נתונים והסתרת פרטי שגיאות ספק הם נושאי הקשחה להמשך. אין בספר טענה שהם הוטמעו או נבדקו במימוש הנוכחי.')
h('12 משאבים הנדרשים לפרויקט')
h('12.1 שעות עבודה ותחומי אחריות',2)
p('הערכת הזמן באפיון המקורי היא 500 שעות. הטבלה מחלקת הערכה זו לתחומים; היא אינה יומן שעות שבוצעו בפועל. בספר מופיע מגיש אחד, ג׳ורדן סיגל, ולכן לא מיוחסת עבודה לחברי צוות שלא צוינו.')
table(['תחום','שעות משוערות'],[('אפיון ותכנון',50),('מובייל וממשק משתמש',130),('שרת ולוגיקה עסקית',100),('מסד נתונים ומיגרציות',45),('AI ומשחקיות',65),('אתר ניהול',50),('בדיקות ותיקונים',35),('תיעוד והכנה להגנה',25),('סה״כ',500)],[12,5])
h('12.2 ציוד נדרש',2)
p('מחשב פיתוח עם חיבור לאינטרנט ושטח להרצת סביבות העבודה, מסד נתונים ואמולטור לפי הצורך. לבדיקה מעשית של הרשאות, התראות ומשחק נדרש מכשיר Android או iOS תואם. אין דרישה לציוד חישה ייעודי מעבר ליכולות הטלפון.')
h('12.3 תוכנות נדרשות',2)
bullets(['Visual Studio או Visual Studio Code ו־.NET SDK 10 לפיתוח השרת.','Node.js ומנהל חבילות להתקנת תלויות React ו־Expo.','SQL Server וכלי ניהול מתאים לבדיקת נתונים ומיגרציות.','Expo Go תואם או בניית פיתוח, ודפדפן עדכני לאתר הניהול.','Git לניהול גרסאות ו־Word לעריכת ספר הפרויקט.'])
h('12.4 ידע חדש שנדרש ללמוד',2)
bullets(['פיתוח מסכים וניווט ב־React Native וב־Expo וטיפול בהרשאות מכשיר.','בניית REST API ב־ASP.NET Core, אימות JWT והרשאות לפי תפקיד ובעלות.','מיפוי נתונים ב־EF Core, מיגרציות, אינדקסים ועדכונים מקבילים.','שילוב שירות AI, הכנת הקשר נתונים ובדיקת טיוטות לפני פרסום.','תכנון XP, רמות, משימות ופעולות שאינן מעניקות פרס כפול.','כתיבת בדיקות אוטומטיות והפרדה בין תוצאות בדיקה מקומית לבין פריסה בפועל.'])
h('12.5 ספרות ומקורות מידע',2)
p('מקורות ראשוניים לפיתוח ולתיאור הטכנולוגיות, לצד קובצי הפרויקט שעליהם מבוססים פרקי המימוש:')
for label,url in [('ASP.NET Core','https://learn.microsoft.com/en-us/aspnet/core/'),('Entity Framework Core','https://learn.microsoft.com/en-us/ef/core/'),('Expo','https://docs.expo.dev/'),('Groq','https://console.groq.com/docs/overview'),('הודעת Intuit על העברת Mint ל־Credit Karma','https://www.creditkarma.com/about/releases/intuit-credit-karma-welcomes-all-minters')]:
    p(label);p(url,ltr=True)
paras('קוד הפרויקט: תיקיות BudgetBounder.Api, BudgetBounderMobile ו־BudgetBounder-web. מודלים ובקרים הם המקור לשמות הישויות, חוזי הנתונים והכללים העסקיים.',
'תיעוד מקומי: docs/requirements-audit-2026-09-05.md, docs/admin-update-running.md ו־docs/admin-update-migrations.sql. קובצי הבדיקות במאגר משלימים את התיעוד.',
'מסמך הדוגמה שסופק שימש כמודל למבנה הפרקים, לפירוט הטכנולוגי ולטבלאות הבדיקה. תוכן אפליקציית ההיכרויות שבו אינו חלק מ־BudgetBounder.')
h('13 תכנית עבודה ושלבים למימוש הפרויקט')
p('מועדי התכנית המקורית נשמרו. אלה מועדי יעד, ולא אישור שכל שלב הושלם במועד.')
table(['שלב','מועד יעד','תוצר'],[('אפיון','19/06/2026','דרישות, תרחישים ומודל נתונים'),('עיצוב וקידוד','31/08/2026','מובייל, API, מסד, משחק ואתר ניהול'),('בדיקות ותיקונים','09/2026','בדיקות אוטומטיות ותהליכיות ותיקון תקלות'),('הגנה פנימית','09/2026','הדגמה והסבר טכני'),('הגשה','30/09/2026','ספר פרויקט וקוד להגשה'),('הגנה חיצונית','10/2026','הצגת המערכת ומענה לשאלות')],[4,4,9])
p('סדר תלות מומלץ למימוש: תחילה משתמשים ומסד, לאחר מכן תנועות ותקציב, יעדים ומשימות, חישובי התקדמות, משחק ו־AI, ולבסוף כלי ניהול ובדיקות שילוב. שינוי סכמה מחייב עדכון API לפני הפעלת לקוח המסתמך עליה.')
h('14 בדיקות')
p('הפרק מפריד בין תרחישי קבלה לבין תוצאות מתועדות. תוצאות הריצה המובאות בהמשך נלקחו מדוח המאגר מ־05/09/2026; לא בוצעה ריצת בדיקות חדשה לצורך עריכת הספר. אין לראות בתוצאה מקומית אישור לעבודה מול שרת ייצור או מכשיר פיזי.')
h('14.1 בדיקות תהליכיות',2)
table(['תהליך','צעדי בדיקה','תוצאה רצויה'],[
('חשבון','הרשמה, התחברות ושליפת פרופיל','JWT תקין ונתוני המשתמש בלבד; דואר קיים נדחה'),
('הוצאה ותקציב','הגדרת תקציב והוספת הוצאה עם Auto','התנועה נשמרת, מסווגת והסיכום החודשי מתעדכן'),
('הכנסה','הוספת Income וצפייה בסיכום','הכנסה מוצגת בנפרד ואינה מצטרפת לסכום הוצאות'),
('חיסכון','יצירת יעד והפקדה','סכום היעד ומדדי השבוע והחודש מתעדכנים'),
('משימת AI','יצירת טיוטה, סקירה ואישור מנהל','טיוטה אינה זמינה לביצוע לפני אישור'),
('תגמול משימה','השלמת תנאים ושליחת השלמה פעמיים','XP ניתן פעם אחת בלבד'),
('תוצאת משחק','שליחת תוצאה תקינה ושליחה חוזרת באותו מזהה','רשומה אחת ופרס אחד'),
('תוצאה לא תקינה','שליחת משך או ניקוד מחוץ לטווח','Flagged ללא XP'),
('הרשאות','ניסיון של משתמש לגשת לנתוני משתמש אחר או לניהול','גישה נדחית'),
('הקפאה','מנהל מקפיא משתמש והמשתמש שולח בקשה עם JWT קיים','הבקשה נדחית'),
('דוחות','סינון תאריכים וייצוא CSV','הייצוא מתאים לדוח המסונן ולתחום הנתונים'),
('מכשיר ורשת','סירוב הרשאה, כשל רשת וחזרה למסך','הודעה ברורה ללא הצגת הצלחה כוזבת')],[3,7,7])
p('מצב אימות תהליכי: דוח המאגר מתעד בדיקות ממשק עם נתוני בדיקה מקומיים לפתיחת ושליחת טופס משימה, הצגת אנליטיקה, הורדת CSV ותצוגה ברוחב 390 פיקסלים. תהליכים עם ספק AI חי, מסד מרוחק וטלפון פיזי עדיין דורשים אימות בסביבת ההגשה. לא מולאה תוצאה שהתקבלה עבור תרחיש שלא תועד כבדיקה שבוצעה.')
h('14.2 בדיקות יחידה',2)
table(['רכיב','מקרי בדיקה','תוצאה צפויה'],[
('ProgressionService','XP לפני סף, בסף ומעליו; פעילות חוזרת באותו יום','רמה לפי הסף ורצף ללא ספירה כפולה'),
('FinancialSummaryService','תחילת וסוף חודש, Income ו־Expense','רק תנועות החודש; הכנסות והוצאות בנפרד'),
('TransactionCategorizer','מילות מפתח בעברית ובאנגלית וטקסט לא מוכר','קטגוריה מתאימה או Other'),
('GameResultValidator','גבולות זמן וניקוד, חישוב XP','תוצאה תקינה בטווח; חריגה מסומנת'),
('Mission administration','טיוטה, אישור, דחייה ומשימה שכבר התקדמה','שמירה על סטטוס ועל התקדמות קיימת'),
('Savings / achievements','הפקדות מתוארכות ופתיחת תגמול חוזרת','סיכום לפי תאריך וללא כפילות'),
('Mobile helpers','כללי תזכורות, כתובת API ותוצאת Tower','התנהגות עקבית בקלטי גבול'),
('Admin helpers','נתוני אנליטיקה, אסימון וכתובת API','עיבוד וייצוא בהתאם לנתונים')],[5,7,5])
table(['חלק במערכת','תוצאה שתועדה ב־05/09/2026'],[('API','52 בדיקות עברו; כולל שירותים ובקרים'),('מובייל','54 בדיקות עברו; TypeScript ו־ESLint עברו; יוצאו חבילות Android ו־iOS'),('אתר מנהל','11 בדיקות עברו; בניית TypeScript/Vite ו־ESLint עברו'),('מסד נתונים','לא נמצאו שינויים במודל ללא מיגרציה; נוצר סקריפט SQL אידמפוטנטי')],[5,12])
p('בדיקות בקרים ושירותים אינן כולן בדיקות יחידה מבודדות. לצורך שחזור מריצים את פקודות הבדיקה בתיקיית הרכיב המתאים לאחר התקנת תלויות, ומתעדים מחדש תאריך, מספר בדיקות וגרסת קוד. אין להסתמך על מספרי הדוח עבור שינויים שנעשו לאחריו.')
for cmd in ['dotnet test BudgetBounder.Api.Tests/BudgetBounder.Api.Tests.csproj','npm test','npm run lint','npm run build']:p(cmd,ltr=True)
h('15 בקרת גרסאות')
paras('הפרויקט מנוהל ב־Git. כל שינוי צריך להיכלל בקומיט בעל תיאור ברור, עם בדיקות המתאימות להיקפו. קובצי נעילת החבילות נשמרים כדי לאפשר התקנה עקבית, ומיגרציות נשמרות לצד קוד השרת כדי לקשר בין גרסת תוכנה לסכמת נתונים.',
'היסטוריית המאגר שנבדקה כוללת את הרשומות הבאות. תיאור הקומיט מובא כפי שנרשם; כותרת קצרה אינה מספיקה להסיק מה היה היקף השינוי.')
table(['מזהה','תאריך','תיאור רשום'],[('5864552','06/09/2026','d'),('bda52ab','05/09/2026','fix: resync lockfiles for admin web and mobile'),('f3f9167','05/09/2026','Merge pull request #2: codex/admin-missions-analytics')],[3,4,10])
p('מהדורת הספר הנוכחית מתוארכת 28/09/2026 וכוללת השלמת פרקים והתאמה למימוש שנבדק. מספר הגרסה בחבילת המובייל הוא 1.0.0 ובחבילת האתר 0.0.0; אלה ערכי metadata של החבילות ואינם מעידים על פריסה. לשחרור להגשה מומלץ ליצור תג גרסה לאחר הרצת הבדיקות ולצרף את הספר, הוראות ההרצה והמיגרציות המתאימות.')
h('16 קטעי קוד')
p('הקטעים הבאים נלקחו מקובצי המקור של הפרויקט. הם מדגימים את הכללים המרכזיים ואינם כוללים מפתחות, סיסמאות או נתוני משתמשים.')
def codeblock(text):
    d.paragraphs[-1].paragraph_format.keep_with_next=True
    text=text.replace(' || ', ' ||\n            ').replace(' + request.', ' +\n            request.')
    lines=text.splitlines()
    for idx,line in enumerate(lines):
        a=p(line,ltr=True);a.paragraph_format.space_after=Pt(0);a.paragraph_format.line_spacing=1
        a.paragraph_format.keep_with_next=idx<len(lines)-1
        for r in a.runs:r.font.name='Consolas';r.font.size=Pt(8)
    p()
h('16.1 חישוב רמה לפי XP',2)
p('מקור: Services/ProgressionService.cs. התאמת הטווחים נעשית מהסף הגבוה לנמוך, ולכן מתקבלת הרמה הגבוהה ביותר שהמשתמש זכאי לה. אותה פונקציה משמשת גם לאחר הענקת XP כדי למנוע הבדל בין הניקוד לרמה.')
source=(ROOT/'BudgetBounder/BudgetBounder.API/BudgetBounder.Api/Services/ProgressionService.cs').read_text(encoding='utf-8-sig')
codeblock(source[source.index('    public static int CalculateLevel'):source.index('    public static void UpdateStreak')].strip())
h('16.2 אימות וחישוב פרס משחק',2)
p('מקור: Services/GameResultValidator.cs. הפונקציה דורשת מזהה תוצאה, בודקת טווחי זמן וערכים ורק לאחר מכן מחשבת פרס. הבדיקה אינה פונה למסד; מניעת הענקה כפולה נעשית בבקר ובאינדקס הייחודי.')
source=(ROOT/'BudgetBounder/BudgetBounder.API/BudgetBounder.Api/Services/GameResultValidator.cs').read_text(encoding='utf-8-sig')
codeblock(source[source.index('    public static GameResultValidation Validate'):].rstrip().removesuffix('}').strip())
h('16.3 סיכום הכנסות והוצאות',2)
p('מקור: Services/FinancialSummaryService.cs. הקטע פועל על תנועות שכבר סוננו לפי משתמש וחודש. הוא מפריד בין סוגי תנועה ומקבץ הוצאות לפי קטגוריה, כדי שניתן יהיה להציג גם סכום כולל וגם התפלגות.')
source=(ROOT/'BudgetBounder/BudgetBounder.API/BudgetBounder.Api/Services/FinancialSummaryService.cs').read_text(encoding='utf-8-sig')
codeblock(source[source.index('        var income ='):source.index('        return new MonthlyFinancialSummary')].strip())
p('החישוב מחזיר יחד עם הסכומים את budget - spent. עבור n תנועות שנשלפו, המעבר לסכימה ולקיבוץ הוא בקירוב ליניארי ב־n; עלות השאילתה תלויה גם באינדקסים ובכמות הנתונים במסד.')

# Preserve the source type system: David for Hebrew complex script and
# Times New Roman for Latin text, with the original 14-point Hebrew body size.
for a in d.paragraphs:
    for r in a.runs:
        if r.font.name=='Consolas':continue
        rp=r._r.get_or_add_rPr();fonts=rp.get_or_add_rFonts()
        fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');fonts.set(qn('w:cs'),'David')
        if re.search('[\u0590-\u05ff]',r.text):
            r.font.rtl=True
            sz=rp.find(qn('w:szCs'))
            if sz is None:sz=OxmlElement('w:szCs');rp.append(sz)
            sz.set(qn('w:val'),'28')
d.core_properties.title='ספר פרויקט גמר BudgetBounder'
d.core_properties.subject='תיעוד מערכת BudgetBounder'
d.core_properties.author='ג׳ורדן סיגל'
d.save(OUT)
print(OUT)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
