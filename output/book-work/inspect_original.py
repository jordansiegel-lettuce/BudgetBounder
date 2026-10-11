from docx import Document
from docx.oxml.ns import qn
from collections import Counter
files=['C:/Users/jordi_vyuby09/Desktop/school/final proj docs/ספר פרויקט גמר.docx','C:/Users/jordi_vyuby09/Desktop/school/final proj docs/ספר פרויקט גמר - BudgetBounder - מלא.docx']
for f in files:
 d=Document(f);print('FILE',f)
 inds=[6,13,20,31,113,114,115,116,117,119,126,128,129,153,154,253,339,436,437,438,469,545] if len(d.paragraphs)>400 else [0,1,4,15,16,17,18,19,20]
 for i in inds:
  p=d.paragraphs[i]
  print(i,p.text[:100],'STYLE',p.style.name)
  print('PPR',[(x.tag.split('}')[-1],dict(x.attrib)) for x in p._p.pPr] if p._p.pPr is not None else None)
  print('RUNS',[(r.text[:25],r.font.name,r.font.size.pt if r.font.size else None,r.bold,r.font.rtl,dict(r._r.rPr.rFonts.attrib) if r._r.rPr is not None and r._r.rPr.rFonts is not None else None) for r in p.runs][:5])
 print('STYLES')
 for name in ['Normal','Heading 1','Heading 2','List Paragraph']:
  s=d.styles[name];print(name,s.element.xml[s.element.xml.index('>')+1:])
 print('FONTS',Counter((r.font.name,r.font.size.pt if r.font.size else None) for p in d.paragraphs for r in p.runs if r.text.strip()))
 print('SECTIONS',[(s.page_width.cm,s.page_height.cm,s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm) for s in d.sections])
