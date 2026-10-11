from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from pathlib import Path

f=Path('C:/Users/jordi_vyuby09/Desktop/school/final proj docs/ספר פרויקט גמר - BudgetBounder - מלא - עיצוב מקורי.docx')
d=Document(f)
ps=[p.text.strip() for p in d.paragraphs[113:]]
heads=[]
start=0
for i in range(1,17):
    target=f'{i}. '
    pos=next(j for j in range(start,len(ps)) if ps[j].startswith(target))
    heads.append((i,pos,ps[pos]))
    start=pos+1
assert len(d.tables)==26 and len(d.inline_shapes)==5
s=d.styles['Normal'];sec=d.sections[0]
print('chapters',heads)
sz=s.element.rPr.find(qn('w:szCs'))
print('font',dict(s.element.rPr.rFonts.attrib),'cs-size',sz.get(qn('w:val')) if sz is not None else 'inherited')
print('margins',sec.top_margin.cm,sec.bottom_margin.cm,sec.left_margin.cm,sec.right_margin.cm)
print('pages',len(PdfReader('output/book-work/original-style-final.pdf').pages),'bytes',f.stat().st_size)
